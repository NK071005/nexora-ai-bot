"""Generate the InternFlow sample spreadsheets."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples"

HEADERS = [
    "Full Name",
    "Email",
    "Role",
    "Department",
    "Company Name",
    "Start Date",
    "End Date",
    "Duration Weeks",
    "Timezone",
    "Initial Task",
]

CLEAN_ROWS = [
    ["Aditi Verma", "aditi.verma@example.com", "Research Intern", "Research", "Nexora Labs", "2026-11-02", "2027-01-22", 12, "Asia/Kolkata", "market-research-brief"],
    ["Kabir Singh", "kabir.singh@example.com", "Product Intern", "Product", "Nexora Labs", "2026-11-02", "2027-01-22", 12, "Asia/Kolkata", "market-research-brief"],
    ["Neha Kulkarni", "neha.kulkarni@example.com", "Data Intern", "Analytics", "Nexora Labs", "2026-11-09", "2027-01-29", 12, "Asia/Kolkata", "squad-analytics-review"],
    ["Arjun Menon", "arjun.menon@example.com", "Brand Intern", "Brand", "Nexora Labs", "2026-11-09", "2027-01-29", 12, "Asia/Kolkata", "brand-perception-review"],
    ["Sara Ali", "sara.ali@example.com", "Research Intern", "Research", "Nexora Labs", "2026-11-16", "2027-02-05", 12, "Asia/Kolkata", "market-research-brief"],
    ["Vikram Bose", "vikram.bose@example.com", "Product Intern", "Product", "Nexora Labs", "2026-11-16", "2027-02-05", 12, "Asia/Kolkata", "competitive-positioning-memo"],
    ["Ishita Rao", "ishita.rao@example.com", "Analytics Intern", "Analytics", "Nexora Labs", "2026-11-23", "2027-02-12", 12, "Asia/Kolkata", ""],
    ["Dev Patel", "dev.patel@example.com", "Operations Intern", "Operations", "Nexora Labs", "2026-11-23", "2027-02-12", 12, "Asia/Kolkata", ""],
]

ERROR_ROWS = [
    ["Riya Nair", "riya.nair@example.com", "Research Intern", "Research", "Nexora Labs", "2026-11-02", "2027-01-22", 12, "Asia/Kolkata", ""],
    ["Missing Email", "", "Product Intern", "Product", "Nexora Labs", "2026-11-02", "2027-01-22", 12, "Asia/Kolkata", ""],
    ["Bad Email", "not-an-email", "Data Intern", "Analytics", "Nexora Labs", "2026-11-02", "2027-01-22", 12, "Asia/Kolkata", ""],
    ["Duplicate Row", "riya.nair@example.com", "Brand Intern", "Brand", "Nexora Labs", "2026-11-02", "2027-01-22", 12, "Asia/Kolkata", ""],
    ["Bad Dates", "bad.dates@example.com", "Research Intern", "Research", "Nexora Labs", "2027-02-01", "2026-12-01", 4, "Mars/Olympus", ""],
]


def _style(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1B2A4A")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
    widths = [24, 30, 20, 18, 18, 14, 14, 16, 18, 28]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top")


def _write(path: Path, rows: list[list]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Interns"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    _style(sheet)
    notes = workbook.create_sheet("How to import")
    notes.append(["Step", "Action"])
    notes.append([1, "Sign in as an organization administrator."])
    notes.append([2, "Open InternFlow → Import interns."])
    notes.append([3, "Upload this workbook. A preview validates every row before anything is created."])
    notes.append([4, "Confirm the import. Each intern receives an activation link; email status is shown honestly."])
    notes.append(["", ""])
    notes.append(["Required columns", "Full Name, Email"])
    notes.append(["Optional columns", "Role, Department, Company Name, Start Date, End Date, Duration Weeks, Timezone, Initial Task"])
    notes.append(["Initial Task", "Accepts a task slug (market-research-brief), an id, or a sequence number."])
    _style(notes)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


if __name__ == "__main__":
    _write(SAMPLES / "internflow_interns.xlsx", CLEAN_ROWS)
    _write(SAMPLES / "internflow_interns_with_errors.xlsx", ERROR_ROWS)
    print(f"Wrote {SAMPLES / 'internflow_interns.xlsx'}")
    print(f"Wrote {SAMPLES / 'internflow_interns_with_errors.xlsx'}")
